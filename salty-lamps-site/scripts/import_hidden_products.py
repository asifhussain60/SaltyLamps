#!/usr/bin/env python3
"""Creates the hidden products the owner listed in the return workbook.

The public preview cannot show a product that is hidden from customers, and the
database that does know them belongs to the retired account. The owner lists them
on the "Hidden products" tab of the workbook built by build_owner_return_workbook.py
and this script creates them through the admin API, so every change lands in the
audit log exactly like a human edit.

WHAT IT WILL AND WILL NOT DO
  - Reports first; writes only with --apply.
  - Creates products. It never edits or deletes one. A product whose name or slug is
    already in the shop is skipped and named, because the shop is the newer truth.
  - A product stays hidden unless its "Put on sale now?" cell says Yes, and that needs
    a stock count. A blank stock count creates the option with 0 in stock, never a
    guess.
  - Refuses a product code already used by another product, the defect this whole
    exercise exists to avoid.
  - Retries are safe: the request id is derived from the product's content, so a
    lost response replays instead of duplicating.
  - Reads everything back after writing and says exactly how many matched.

USAGE
    python3 scripts/import_hidden_products.py <file.xlsx> [--apply] [--api=<url>]

Against a deployed shop set CF_ACCESS_CLIENT_ID and CF_ACCESS_CLIENT_SECRET in the
environment. They are passed as headers and never printed.
"""

import hashlib
import json
import re
import sys
import urllib.error
import urllib.request
from decimal import Decimal, InvalidOperation
from importlib.machinery import SourceFileLoader
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
owner = SourceFileLoader("import_owner_workbook", str(HERE / "import_owner_workbook.py")).load_module()

API = "http://localhost:8788"
SHEET = "Hidden products"

# Header text is matched by prefix, so the explanatory second line can change.
PREFIXES = {
    "name": "product name", "category": "category", "description": "description",
    "size": "size / option", "code": "product code", "price": "price",
    "stock": "stock you have now", "packed_weight": "packed shipping weight",
    "product_weight_min": "product weight from", "product_weight_max": "product weight to",
    "on_sale": "put on sale now", "note": "notes",
}


def norm(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def slugify(value):
    return re.sub(r"(^-|-$)", "", re.sub(r"[^a-z0-9]+", "-", norm(value).lower()))


def base_code(code):
    return re.sub(r"\s*\(\d+\)\s*$", "", str(code)).strip().casefold()


def read_rows(ws):
    """Return (list of raw row dicts, errors). Blank rows are skipped."""
    header = {}
    for r in range(1, min(ws.max_row, 10) + 1):
        cells = [norm(c.value).lower() for c in ws[r]]
        found = {key: i for key, prefix in PREFIXES.items()
                 for i, text in enumerate(cells, 1) if text.startswith(prefix)}
        if "name" in found and "code" in found:
            header, first = found, r + 1
            break
    if not header:
        return [], ["Cannot find the Hidden products header (Product name and Product code are required)."]
    rows = []
    for r in range(first, ws.max_row + 1):
        values = {key: ws.cell(r, col).value for key, col in header.items()}
        if all(v in (None, "") for v in values.values()):
            continue
        values["row"] = r
        rows.append(values)
    return rows, []


def parse_price(value):
    try:
        amount = Decimal(str(value).strip().lstrip("£"))
    except (InvalidOperation, ValueError):
        return None
    pence = amount * 100
    if amount <= 0 or pence != pence.to_integral_value():
        return None
    return int(pence)


def group_products(rows, category_slugs):
    """Group rows into products and validate them. Returns (products, problems)."""
    problems, products = [], {}
    for row in rows:
        at = f"row {row['row']}"
        name = norm(row["name"])
        if not name:
            problems.append(f"{at}: has details but no product name.")
            continue
        product = products.setdefault(name.casefold(), {
            "name": name, "description": "", "category": "", "on_sale": False, "options": [], "rows": []})
        product["rows"].append(row["row"])
        for key in ("description", "category"):
            given = norm(row[key])
            if given and product[key] and given != product[key]:
                problems.append(f"{at}: {key} differs from an earlier row of \"{name}\".")
            elif given:
                product[key] = given
        if norm(row["on_sale"]).lower() == "yes":
            product["on_sale"] = True
        elif norm(row["on_sale"]) and norm(row["on_sale"]).lower() != "no":
            problems.append(f"{at}: \"Put on sale now?\" must be Yes or left blank.")
        code, label = norm(row["code"]), norm(row["size"])
        if not code:
            problems.append(f"{at}: \"{name}\" needs a product code.")
        pence = parse_price(row["price"])
        if pence is None:
            problems.append(f"{at}: \"{name}\" needs a price in pounds and pence, more than zero.")
        stock = None
        if isinstance(row["stock"], float) and row["stock"].is_integer():
            row["stock"] = int(row["stock"])  # Excel may hand back 5.0 for a typed 5
        if row["stock"] not in (None, ""):
            if isinstance(row["stock"], bool) or not re.fullmatch(r"\d+", str(row["stock"]).strip()):
                problems.append(f"{at}: stock must be a whole number, 0 or more.")
            else:
                stock = int(str(row["stock"]).strip())
        try:
            weights = owner.weight_updates({
                "packed_weight": row["packed_weight"], "product_weight_min": row["product_weight_min"],
                "product_weight_max": row["product_weight_max"]})
            lo, hi, packed = (weights.get(k) for k in ("product_weight_min_g", "product_weight_max_g", "packed_weight_g"))
            if (lo is None) != (hi is None) or (lo is not None and lo > hi):
                raise ValueError("Enter both ends of a valid product weight range.")
            if packed is not None and hi is not None and packed < hi:
                raise ValueError("Packed weight must cover the maximum product weight.")
        except ValueError as e:
            problems.append(f"{at}: {e}")
            weights = {}
        product["options"].append({"code": code, "label": label, "pence": pence, "stock": stock,
                                   "weights": weights, "row": row["row"]})
    for product in products.values():
        name = product["name"]
        if len(product["options"]) > 1 and any(not o["label"] for o in product["options"]):
            problems.append(f"\"{name}\": every size needs a Size / option when there is more than one row.")
        seen = set()
        for o in product["options"]:
            if (o["code"].casefold(), o["label"].casefold()) in seen:
                problems.append(f"row {o['row']}: \"{name}\" lists {o['code']} / {o['label'] or '(no size)'} twice.")
            seen.add((o["code"].casefold(), o["label"].casefold()))
        if product["on_sale"] and any(o["stock"] is None for o in product["options"]):
            problems.append(f"\"{name}\": it can't go on sale until every size has a stock count.")
        if product["category"] and product["category"].casefold() not in category_slugs:
            problems.append(f"\"{name}\": category \"{product['category']}\" is not one of the shop's categories.")
    # Codes may repeat within one product (pack sizes) but never across products.
    owners = {}
    for product in products.values():
        for o in product["options"]:
            if o["code"]:
                owners.setdefault(base_code(o["code"]), set()).add(product["name"])
    for code, names in sorted(owners.items()):
        if len(names) > 1:
            problems.append(f"code {code!r} is used by more than one product: " + "; ".join(sorted(names)))
    return list(products.values()), problems


def plan(products, live, category_slugs):
    """Split into products to create and products already in the shop; add clash problems."""
    names = {p["name"].casefold(): p for p in live}
    slugs = {p.get("slug"): p for p in live}
    live_codes = {}
    for p in live:
        for s in p.get("skus", []):
            live_codes.setdefault(base_code(s["sku"]), set()).add(p["name"])
    create, existing, problems = [], [], []
    for product in products:
        hit = names.get(product["name"].casefold()) or slugs.get(slugify(product["name"]))
        if hit:
            existing.append((product, hit))
            continue
        for o in product["options"]:
            clash = live_codes.get(base_code(o["code"]), set())
            if clash:
                problems.append(f"row {o['row']}: code {o['code']!r} is already used in the shop by "
                                + "; ".join(sorted(clash)))
        create.append(product)
    return create, existing, problems


def payload(product, category_slugs):
    categories = [s for s in ("all-products", category_slugs.get(product["category"].casefold())) if s]
    options = []
    for o in product["options"]:
        option = {"sku": o["code"], "variant_label": o["label"], "price_pence": o["pence"],
                  "track_mode": "quantity", "quantity": o["stock"] if o["stock"] is not None else 0}
        option.update(o["weights"])
        options.append(option)
    body = {"product": {"name": product["name"], "slug": slugify(product["name"]),
                        "description": product["description"], "categories": ",".join(categories),
                        "visible": 1 if product["on_sale"] else 0},
            "skus": options}
    digest = hashlib.sha256(json.dumps(body, sort_keys=True).encode()).hexdigest()[:40]
    body["requestId"] = f"hidden_{digest}"
    return body


def get(path):
    return owner.fetch(path)


def post(path, body):
    req = urllib.request.Request(f"{API}{path}", method="POST", data=json.dumps(body).encode(),
                                 headers={**owner.HEADERS, "content-type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return True, json.load(r)
    except urllib.error.HTTPError as e:
        return False, e.read().decode()[:300]


def read_back(product, live):
    """True when the shop now holds this product with the values that were sent."""
    hit = next((p for p in live if p["name"].casefold() == product["name"].casefold()), None)
    if hit is None or bool(hit.get("visible", 1)) != product["on_sale"]:
        return False
    have = {(s["sku"], s["variant_label"]): s for s in hit.get("skus", [])}
    for o in product["options"]:
        s = have.get((o["code"], o["label"]))
        if s is None or s["price_pence"] != o["pence"] or s.get("quantity") != (o["stock"] or 0):
            return False
        if any(s.get(k) != v for k, v in o["weights"].items() if k.endswith("_g")):
            return False
    return len(have) == len(product["options"])


def main():
    global API
    argv = sys.argv[1:]
    positional = [a for a in argv if not a.startswith("--")]
    if not positional:
        sys.exit("usage: import_hidden_products.py <file.xlsx> [--apply] [--api=<url>]")
    apply = "--apply" in argv
    api = next((a.split("=", 1)[1] for a in argv if a.startswith("--api=")), None)
    API = (api or API).rstrip("/")
    owner.API = API
    path = Path(positional[0])
    if not path.exists():
        sys.exit(f"No such file: {path}")
    book = load_workbook(path, data_only=False)
    if SHEET not in book.sheetnames:
        sys.exit(f"No \"{SHEET}\" tab in {path.name}.")
    rows, errors = read_rows(book[SHEET])
    if errors:
        sys.exit("\n".join(errors))

    taxonomy = get("/api/categories")["categories"]
    slugs = {c["name"].casefold(): c["slug"] for c in taxonomy}
    products, problems = group_products(rows, slugs)
    live = get("/api/admin/products")["products"]
    create, existing, clashes = plan(products, live, slugs)
    problems.extend(clashes)

    print(f"\n=== {path.name} ({SHEET}) → {API} ===\n")
    print(f"  rows read                 : {len(rows)}")
    print(f"  products to create        : {len(create)}")
    print(f"  already in the shop, skipped: {len(existing)}")
    for product in create:
        state = "ON SALE" if product["on_sale"] else "hidden"
        print(f"\n  + {product['name']}  [{state}]")
        for o in product["options"]:
            stock = "stock not given → 0" if o["stock"] is None else f"stock {o['stock']}"
            packed = o["weights"].get("packed_weight_g")
            weight = f"packed {packed / 1000:g} kg" if packed else "packed weight not given"
            print(f"      {o['code']:<12} {o['label'] or '(one size)':<14} £{o['pence'] / 100:>7.2f}  {stock}, {weight}")
    for product, hit in existing:
        print(f"\n  = {product['name']}: already in the shop as \"{hit['name']}\". Not changed here; "
              "edit it in the admin, or on the Public products tab if its options are already there.")
    if problems:
        print(f"\n  !! {len(problems)} problem(s) — nothing will be written until these are resolved:")
        for p in problems[:30]:
            print(f"     {p}")
        print("\nRefusing to write while there are problems above. Fix the sheet and re-run.\n")
        sys.exit(1)
    if not apply:
        print("\nReport only. Re-run with --apply to create these through the admin API.\n")
        return
    made = 0
    for product in create:
        ok, result = post("/api/admin/products/save", payload(product, slugs))
        made += ok
        if not ok:
            print(f"  FAILED {product['name']}: {result}")
    after = get("/api/admin/products")["products"]
    verified = sum(read_back(p, after) for p in create)
    print(f"\n  {made} of {len(create)} product(s) written; {verified} of {len(create)} read back "
          "exactly as sent.\n")
    if verified != len(create):
        sys.exit(1)


if __name__ == "__main__":
    main()
